"""Generate fictional PDF and DOCX resumes from the adjacent Markdown source."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "jordan-lee-resume.md"
PDF_OUTPUT = ROOT / "jordan-lee-resume.pdf"
DOCX_OUTPUT = ROOT / "jordan-lee-resume.docx"

# compact_reference_guide with named A4 resume override:
# A4 portrait, 18 mm margins, Arial 10.5 pt, 1.15 line spacing,
# restrained blue hierarchy, no tables, running headers, or footers.
BLUE = "2457A6"
DARK = "15233B"
MUTED = "53627A"


def source_lines() -> list[str]:
    return [line.rstrip() for line in SOURCE.read_text(encoding="utf-8").splitlines()]


def set_run_font(run, *, size: float, color: str = DARK, bold: bool = False) -> None:
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def configure_docx_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, before, after in (
        ("Heading 1", 14, 10, 4),
        ("Heading 2", 11.5, 7, 2),
    ):
        style = document.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLUE)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    bullet = document.styles.add_style("Resume Bullet", WD_STYLE_TYPE.PARAGRAPH)
    bullet.base_style = document.styles["List Bullet"]
    bullet.font.name = "Arial"
    bullet.font.size = Pt(10.5)
    bullet.paragraph_format.left_indent = Mm(5.3)
    bullet.paragraph_format.first_line_indent = Mm(-5.3)
    bullet.paragraph_format.space_after = Pt(3)
    bullet.paragraph_format.line_spacing = 1.15


def generate_docx() -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(18)
    section.bottom_margin = Mm(18)
    section.left_margin = Mm(18)
    section.right_margin = Mm(18)
    section.header_distance = Mm(8)
    section.footer_distance = Mm(8)
    configure_docx_styles(document)

    lines = source_lines()
    for index, line in enumerate(lines):
        if not line:
            continue
        if line.startswith("# "):
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(2)
            set_run_font(paragraph.add_run(line[2:]), size=22, color=DARK, bold=True)
        elif line.startswith("## "):
            document.add_paragraph(line[3:], style="Heading 1")
        elif line.startswith("### "):
            document.add_paragraph(line[4:], style="Heading 2")
        elif line.startswith("- "):
            document.add_paragraph(line[2:], style="Resume Bullet")
        elif index in (2, 4):
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(5 if index == 4 else 1)
            set_run_font(
                paragraph.add_run(line),
                size=11 if index == 2 else 9.5,
                color=BLUE if index == 2 else MUTED,
                bold=index == 2,
            )
        else:
            document.add_paragraph(line)

    document.core_properties.title = "Jordan Lee - Fictional Resume"
    document.core_properties.author = "Job Copilot Fixture Generator"
    document.core_properties.subject = "Synthetic resume fixture for local validation"
    document.save(DOCX_OUTPUT)


def generate_pdf() -> None:
    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "ResumeBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=10.5,
        leading=12.1,
        textColor=HexColor(f"#{DARK}"),
        spaceAfter=4,
    )
    title = ParagraphStyle(
        "ResumeTitle",
        parent=body,
        fontName="Helvetica-Bold",
        fontSize=22,
        leading=25,
        alignment=TA_CENTER,
        spaceAfter=2,
    )
    subtitle = ParagraphStyle(
        "ResumeSubtitle",
        parent=body,
        fontName="Helvetica-Bold",
        fontSize=11,
        textColor=HexColor(f"#{BLUE}"),
        alignment=TA_CENTER,
        spaceAfter=2,
    )
    contact = ParagraphStyle(
        "ResumeContact",
        parent=body,
        fontSize=9.5,
        textColor=HexColor(f"#{MUTED}"),
        alignment=TA_CENTER,
        spaceAfter=6,
    )
    heading = ParagraphStyle(
        "ResumeHeading",
        parent=body,
        fontName="Helvetica-Bold",
        fontSize=14,
        textColor=HexColor(f"#{BLUE}"),
        spaceBefore=8,
        spaceAfter=3,
        keepWithNext=True,
    )
    role = ParagraphStyle(
        "ResumeRole",
        parent=body,
        fontName="Helvetica-Bold",
        fontSize=11.5,
        textColor=HexColor(f"#{BLUE}"),
        spaceBefore=4,
        spaceAfter=2,
        keepWithNext=True,
    )
    bullet = ParagraphStyle(
        "ResumeBullet",
        parent=body,
        leftIndent=5.3 * mm,
        firstLineIndent=-5.3 * mm,
        bulletIndent=0,
        spaceAfter=3,
    )

    story = []
    for index, line in enumerate(source_lines()):
        if not line:
            continue
        if line.startswith("# "):
            story.append(Paragraph(line[2:], title))
        elif line.startswith("## "):
            story.append(Paragraph(line[3:], heading))
        elif line.startswith("### "):
            story.append(Paragraph(line[4:], role))
        elif line.startswith("- "):
            story.append(Paragraph(f"- {line[2:]}", bullet))
        elif index == 2:
            story.append(Paragraph(line, subtitle))
        elif index == 4:
            story.append(Paragraph(line.replace("|", " | "), contact))
        else:
            story.append(Paragraph(line, body))

    document = SimpleDocTemplate(
        str(PDF_OUTPUT),
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title="Jordan Lee - Fictional Resume",
        author="Job Copilot Fixture Generator",
    )
    document.build(story)


if __name__ == "__main__":
    generate_docx()
    generate_pdf()
    print(f"Generated {PDF_OUTPUT.name} and {DOCX_OUTPUT.name}")
